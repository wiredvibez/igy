"""Adds the IGY logo (top left, every slide) and pride rainbows/hearts to the hand-edited deck.

Works on the edited .pptx itself, not on build-deck.js output, so manual edits are kept.
New pictures get auto-start entrance animations appended to each slide's existing timing.

Usage: python add-pride.py <in.pptx> <pride-assets dir> <out.pptx>
"""

import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

import animate

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
W = 13.333
PRIDE = ["E40303", "FF8C00", "FFED00", "008026", "004DFF", "750787"]


def fragment(xml):
    root = etree.fromstring(f'<root xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}">{xml}</root>')
    return list(root)


def add_effects(slide, effects):
    """effects: [(shape, effect, delay_ms, emphasis_or_None)] — appended to the auto-start group."""
    sld = slide._element
    main = sld.find(".//p:cTn[@nodeType='mainSeq']", NS)
    assert main is not None, f"slide {slide.slide_id} has no animation sequence"
    first_click = main.find("p:childTnLst/p:par/p:cTn", NS)
    conds = [c.get("evt") for c in first_click.findall("p:stCondLst/p:cond", NS)]
    assert "onBegin" in conds, "first animation group is not auto-start"
    group = first_click.find("p:childTnLst/p:par/p:cTn/p:childTnLst", NS)
    ids = animate.Ids(max(int(c.get("id")) for c in sld.iter(f"{{{NS['p']}}}cTn")))
    for shape, effect, delay, emph in effects:
        preset, subtype, fn = animate.ENTRANCE[effect]
        xml = animate.effect_par(ids, shape.shape_id, preset, "entr", subtype, 0, delay, fn)
        if emph:
            eff, edelay = emph.split("@")
            xml += animate.effect_par(ids, shape.shape_id, 32, "emph", 0, 1, int(edelay), animate.teeter)
        for el in fragment(xml):
            group.append(el)


def picture(slide, path, x, y, w, h=None, name="", alt=""):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h if h else w))
    pic.name = name
    pic._element.find(".//p:cNvPr", NS).set("descr", alt)
    return pic


def repoint(slide, pic, path):
    """Swap a picture's image but keep the shape (and so its animation)."""
    _, rid = slide.part.get_or_add_image_part(str(path))
    blip = pic._element.find(".//a:blip", NS)
    old = blip.get(f"{{{NS['r']}}}embed")
    blip.set(f"{{{NS['r']}}}embed", rid)
    # python-pptx's drop_rel only counts r:id references, not r:embed, so check blips ourselves.
    if old != rid and not slide._element.xpath(f'.//a:blip[@r:embed="{old}"]'):
        slide.part.drop_rel(old)


def move(shape, x, y, d=None):
    shape.left, shape.top = Inches(x), Inches(y)
    if d:
        shape.width = shape.height = Inches(d)


def pics(slide):
    return [s for s in slide.shapes if s.shape_type == 13]


def shape_num(shape):
    m = re.search(r"(\d+)$", shape.name)
    return int(m.group(1)) if m else 0


def hearts_row(slide, assets, y, start_delay, right=None, d=0.62, step=0.72):
    # Red first on the right (Hebrew reads right to left), violet last on the left.
    width = step * 5 + d
    x0 = (right - width) if right is not None else (W - width) / 2
    out = []
    for i, color in enumerate(PRIDE):
        x = x0 + width - d - i * step
        pic = picture(slide, assets / f"pride-heart-{i + 1}-{color}.png", x, y, d, name=f"pride heart {i + 1}", alt="לב בצבעי הגאווה")
        emph = f"teeter@{start_delay + 6 * 150 + 400}" if i == 5 else None
        out.append((pic, "pop", start_delay + i * 150, emph))
    add_effects(slide, out)


def main(src, assets_dir, dst):
    assets = Path(assets_dir)
    prs = Presentation(src)
    slides = list(prs.slides)
    assert len(slides) == 20, f"expected the 20-slide edited deck, got {len(slides)}"
    by_text = {}
    for i, s in enumerate(slides):
        text = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        by_text[i + 1] = text

    # Counters ("3/8") sit where the logo goes: shift them right of it.
    for s in slides:
        for sh in s.shapes:
            if sh.has_text_frame and re.fullmatch(r"\d/\d", sh.text_frame.text.strip()) and sh.left < Inches(1.6):
                sh.left = Inches(1.7)

    # 1 · cover: move the top-left heart away from the logo, add a row of pride hearts.
    s = slides[0]
    assert "פעולת פתיחה" in by_text[1]
    tl = [p for p in pics(s) if p.left < Inches(1.6) and p.top < Inches(1.6)]
    assert len(tl) == 1
    move(tl[0], 0.65, 2.7)
    hearts_row(s, assets, 6.42, 2100)

    # 2 · "סבב היכרות": a rainbow above the title.
    s = slides[1]
    assert "סבב היכרות" in by_text[2]
    rb = picture(s, assets / "pride-rainbow.png", (W - 3.8) / 2, 0.95, 3.8, 3.8 * 420 / 800, name="pride rainbow", alt="קשת בצבעי הגאווה")
    add_effects(s, [(rb, "pop", 200, "teeter@1300")])

    # 4 · question about the queer community: its rainbow icon becomes a pride rainbow.
    s = slides[3]
    assert "קהילה הגאה" in by_text[4]
    bottom = max(pics(s), key=lambda p: p.top)
    repoint(s, bottom, assets / "pride-rainbow-badge.png")

    # 16 · break: rainbow in the lower right.
    s = slides[15]
    assert "הפסקה" in by_text[16]
    rb = picture(s, assets / "pride-rainbow.png", 9.2, 5.25, 3.5, 3.5 * 420 / 800, name="pride rainbow", alt="קשת בצבעי הגאווה")
    add_effects(s, [(rb, "rise", 1100, "teeter@1900")])

    # 18 · time capsule: a row of pride hearts under the title.
    s = slides[17]
    assert "קפסולת זמן" in by_text[18]
    hearts_row(s, assets, 5.75, 1200)

    # 20 · closing: the six hearts take the pride colors in the order they pop in.
    s = slides[19]
    assert "שנה מעולה" in by_text[20]
    hearts = sorted(pics(s), key=shape_num)
    assert len(hearts) == 6, len(hearts)
    for i, (h, color) in enumerate(zip(hearts, PRIDE)):
        repoint(s, h, assets / f"pride-heart-{i + 1}-{color}.png")
    tl = [p for p in hearts if p.left < Inches(1.6) and p.top < Inches(1.6)]
    for p in tl:
        move(p, 0.85, 1.75)

    # IGY logo, top left, on every slide (static, on top of everything).
    for s in slides:
        picture(s, assets / "igy-logo-badge.png", 0.3, 0.25, 1.2, name="IGY logo", alt="לוגו איגי")

    prs.save(dst)
    print(f"wrote {dst}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
